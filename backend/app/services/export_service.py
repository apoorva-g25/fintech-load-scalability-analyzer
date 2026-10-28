import io
import csv
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.app.models.transaction import Transaction
from backend.app.models.test_run import TestRun

class ExportService:
    @staticmethod
    def export_transactions_csv(transactions: List[Transaction]) -> str:
        """
        Exports a list of transactions to CSV format.
        """
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Headers
        writer.writerow([
            "transaction_id", "test_run_id", "customer_id", "type",
            "source_currency", "dest_currency", "original_amount",
            "exchange_rate", "converted_amount", "fx_fee",
            "status", "fraud_flagged", "fraud_reason",
            "total_duration_ms", "payment_duration_ms", "fraud_duration_ms", "fx_duration_ms",
            "error_message", "completed_at"
        ])
        
        for tx in transactions:
            writer.writerow([
                tx.id, tx.test_run_id, tx.customer_id, tx.transaction_type,
                tx.source_currency, tx.destination_currency, tx.original_amount,
                tx.exchange_rate or "", tx.converted_amount or "", tx.fx_fee or "",
                tx.status, tx.fraud_flagged, tx.fraud_reason or "",
                tx.total_duration_ms, tx.payment_duration_ms, tx.fraud_duration_ms, tx.fx_duration_ms,
                tx.error_message or "", tx.completed_at.isoformat() if tx.completed_at else ""
            ])
            
        return output.getvalue()

    @staticmethod
    def export_benchmark_workbook(test_runs: List[TestRun]) -> io.BytesIO:
        """
        Generates an Excel workbook with detailed benchmark sheets.
        """
        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style helpers
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F2937")

        # Sheet 1: Test Summary
        ws_sum = wb.create_sheet(title="Test Summary")
        ws_sum.append(["FinTech System Load & Scalability Analyzer - Experimental Summary"])
        ws_sum["A1"].font = title_font
        ws_sum.append([])
        ws_sum.append([
            "Test Run ID", "Name", "Status", "Tx Count", "Mode", "Concurrency",
            "Throughput (TPS)", "Avg Latency (ms)", "P95 Latency (ms)", "Success Rate (%)", "Fraud Count"
        ])
        for cell in ws_sum[3]:
            cell.fill = header_fill
            cell.font = header_font

        for tr in test_runs:
            m = tr.metrics or {}
            ws_sum.append([
                tr.id, tr.name, tr.status, tr.transaction_count, tr.processing_mode, tr.concurrency_limit,
                m.get("throughput_tps", 0.0),
                m.get("avg_response_time_ms", 0.0),
                m.get("p95_response_time_ms", 0.0),
                m.get("success_rate_pct", 0.0),
                m.get("fraud_flagged_count", 0)
            ])

        # Sheet 2: FX Analysis
        ws_fx = wb.create_sheet(title="FX Analysis")
        ws_fx.append(["Foreign Exchange Overhead & Latency Analysis"])
        ws_fx["A1"].font = title_font
        ws_fx.append([])
        ws_fx.append([
            "Test Run ID", "Workload Type", "Intl Tx Count", "FX Success", "Avg FX Latency (ms)", "P95 FX Latency (ms)"
        ])
        for cell in ws_fx[3]:
            cell.fill = header_fill
            cell.font = header_font

        for tr in test_runs:
            m = tr.metrics or {}
            fx = m.get("fx_conversions", {})
            ws_fx.append([
                tr.id, tr.workload_type, fx.get("total_intl", 0), fx.get("success_count", 0),
                fx.get("avg_fx_latency_ms", 0.0), fx.get("p95_fx_latency_ms", 0.0)
            ])

        # Auto-adjust column widths
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        return file_stream
