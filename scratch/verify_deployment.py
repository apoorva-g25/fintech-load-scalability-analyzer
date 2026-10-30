"""Verify the application over local HTTP using a disposable database, never demo data."""
import hashlib
import io
import json
import os
from pathlib import Path
import socket
import signal
import subprocess
import sys
import tempfile
import time

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / 'documentation_context' / 'evidence'


def main():
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    original = hashlib.sha256((ROOT / 'fintech_analyzer.db').read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix='fintech-verification-') as temporary:
        with socket.socket() as listener:
            listener.bind(('127.0.0.1', 0))
            port = listener.getsockname()[1]
        env = dict(os.environ, DATABASE_URL='sqlite:///' + (Path(temporary) / 'verify.db').as_posix(),
                   CORS_ORIGINS='["https://frontend.example"]')
        with (EVIDENCE / 'server.txt').open('w', encoding='utf-8') as log:
            process = subprocess.Popen([sys.executable, '-m', 'uvicorn', 'backend.app.main:app',
                                        '--host', '0.0.0.0', '--port', str(port)], cwd=ROOT, env=env,
                                       stdout=log, stderr=subprocess.STDOUT,
                                       creationflags=subprocess.CREATE_NEW_PROCESS_GROUP if os.name == 'nt' else 0)
            try:
                with httpx.Client(base_url=f'http://127.0.0.1:{port}', timeout=20) as client:
                    for _ in range(100):
                        try:
                            if client.get('/api/health').status_code == 200:
                                break
                        except httpx.ConnectError:
                            pass
                        if process.poll() is not None:
                            raise RuntimeError('Server exited; inspect server.txt')
                        time.sleep(0.1)
                    else:
                        raise RuntimeError('Server startup timed out')
                    examples = []

                    def request(method, path, body=None, expected=200):
                        response = client.request(method, path, json=body) if body is not None else client.request(method, path)
                        assert response.status_code == expected, (path, response.status_code, response.text)
                        try:
                            result = response.json()
                        except ValueError:
                            result = {'content_type': response.headers.get('content-type'), 'bytes': len(response.content)}
                        examples.append({'method': method, 'path': path, 'request': body, 'status': response.status_code, 'response': result})
                        return response

                    request('GET', '/')
                    request('GET', '/api/health')
                    schema = request('GET', '/openapi.json').json()
                    (EVIDENCE / 'openapi.json').write_text(json.dumps(schema, indent=2), encoding='utf-8')
                    examples.pop()  # schema stored separately
                    request('GET', '/docs')
                    request('GET', '/api/config')
                    request('PUT', '/api/config', {'base_fx_delay_ms': 8.0})
                    request('GET', '/api/fx/rates')
                    request('PUT', '/api/fx/rates', {'source_currency': 'USD', 'destination_currency': 'INR', 'rate': 83.25, 'is_active': True})
                    ids = []
                    for mode in ('sequential', 'concurrent'):
                        payload = {'name': f'HTTP-Smoke-{mode}', 'transaction_count': 20, 'processing_mode': mode,
                                   'workload_type': 'mixed', 'random_seed': 77, 'concurrency_limit': 10}
                        run_id = request('POST', '/api/test-runs', payload).json()['id']
                        ids.append(run_id)
                        for _ in range(200):
                            run = client.get(f'/api/test-runs/{run_id}').json()
                            if run['status'] in ('COMPLETED', 'FAILED'):
                                break
                            time.sleep(0.05)
                        assert run['status'] == 'COMPLETED', run
                        request('GET', f'/api/test-runs/{run_id}')
                        request('GET', f'/api/test-runs/{run_id}/status')
                        transactions = request('GET', f'/api/test-runs/{run_id}/transactions?page=1&page_size=10').json()
                        assert transactions['total'] == 20 and len(transactions['items']) == 10
                        assert run['metrics']['notification_metrics']['processed_count'] == 20
                        request('GET', f'/api/bottlenecks/{run_id}')
                        csv = request('GET', f'/api/exports/{run_id}')
                        assert len(csv.text.splitlines()) == 21
                    request('GET', '/api/test-runs')
                    request('GET', '/api/dashboard/summary')
                    request('GET', '/api/fx/analytics')
                    request('GET', '/api/analytics')
                    request('POST', '/api/comparisons', {'test_run_ids': ids})
                    workbook = request('GET', '/api/exports/all/excel')
                    assert load_workbook(io.BytesIO(workbook.content)).sheetnames == ['Test Summary', 'FX Analysis']
                    request('GET', '/api/test-runs/missing', expected=404)
                    request('POST', '/api/test-runs', {'transaction_count': 1}, expected=422)
                    request('POST', '/api/comparisons', {'test_run_ids': []}, expected=400)
                    for origin, permitted in [('https://frontend.example', True), ('https://other.example', False)]:
                        response = client.options('/api/test-runs', headers={'Origin': origin, 'Access-Control-Request-Method': 'POST'})
                        assert (response.headers.get('access-control-allow-origin') == origin) == permitted
                    (EVIDENCE / 'http-verification.json').write_text(json.dumps({'date': '2026-10-30', 'scope': 'local HTTP; disposable SQLite; not hosted', 'examples': examples, 'cors_allow_and_deny': 'passed'}, indent=2), encoding='utf-8')
            finally:
                if os.name == 'nt':
                    process.send_signal(signal.CTRL_BREAK_EVENT)
                else:
                    process.terminate()
                process.wait(timeout=15)
    assert hashlib.sha256((ROOT / 'fintech_analyzer.db').read_bytes()).hexdigest() == original
    print('PASS: startup, health, docs, sequential/concurrent runs, retrieval, exports, errors, CORS; original database unchanged')


if __name__ == '__main__':
    main()
